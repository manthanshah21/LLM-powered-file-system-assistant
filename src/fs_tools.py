"""Safe file-system tools for reading and searching resume files."""

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pypdf import PdfReader


BASE_DIR = Path(__file__).resolve().parent.parent
SUPPORTED_READ_EXTENSIONS = {".txt", ".pdf", ".docx"}


def _resolve_path(filepath):
    """Resolve a project-relative path and reject paths outside the project."""
    candidate = Path(filepath)
    if not candidate.is_absolute():
        candidate = BASE_DIR / candidate
    resolved = candidate.resolve()
    try:
        resolved.relative_to(BASE_DIR)
    except ValueError as exc:
        raise ValueError("Path must be inside the project directory.") from exc
    return resolved


def _metadata(path):
    stat = path.stat()
    return {
        "name": path.name,
        "path": str(path.relative_to(BASE_DIR)).replace("\\", "/"),
        "extension": path.suffix.lower(),
        "size": stat.st_size,
        "modified_at": datetime.fromtimestamp(
            stat.st_mtime, tz=timezone.utc
        ).isoformat(),
    }


def _extract_text(path):
    extension = path.suffix.lower()
    if extension == ".txt":
        return path.read_text(encoding="utf-8-sig")
    if extension == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if extension == ".docx":
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(part for part in parts if part)
    raise ValueError(f"Unsupported file type: {extension or '(no extension)'}")


def read_file(filepath: str) -> dict:
    """Read a TXT, PDF, or DOCX file and return its text and metadata."""
    try:
        path = _resolve_path(filepath)
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {filepath}")
        content = _extract_text(path)
        return {"success": True, "content": content, "metadata": _metadata(path)}
    except Exception as exc:
        return {"success": False, "content": "", "metadata": None, "error": str(exc)}


def list_files(directory: str, extension: str = None) -> list:
    """List immediate files in a project directory, optionally by extension."""
    try:
        folder = _resolve_path(directory)
        if not folder.is_dir():
            return []
        normalized_extension = extension.lower() if extension else None
        if normalized_extension and not normalized_extension.startswith("."):
            normalized_extension = f".{normalized_extension}"
        files = [
            _metadata(path)
            for path in sorted(folder.iterdir(), key=lambda item: item.name.lower())
            if path.is_file()
            and (normalized_extension is None or path.suffix.lower() == normalized_extension)
        ]
        return files
    except (OSError, ValueError):
        return []


def write_file(filepath: str, content: str) -> dict:
    """Write UTF-8 text inside the project, creating parent directories."""
    try:
        path = _resolve_path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return {"success": True, "message": f"File written successfully: {filepath}",
                "metadata": _metadata(path)}
    except Exception as exc:
        return {"success": False, "error": str(exc)}


def search_in_file(filepath: str, keyword: str) -> dict:
    """Case-insensitively find keyword occurrences with surrounding text."""
    if not keyword:
        return {"success": False, "match_count": 0, "matches": [],
                "error": "Keyword must not be empty."}
    result = read_file(filepath)
    if not result["success"]:
        return {"success": False, "match_count": 0, "matches": [],
                "error": result["error"]}

    content = result["content"]
    lowered = content.casefold()
    needle = keyword.casefold()
    matches = []
    start_at = 0
    while (index := lowered.find(needle, start_at)) != -1:
        context_start = max(0, index - 80)
        context_end = min(len(content), index + len(keyword) + 80)
        matches.append({
            "position": index,
            "line": content.count("\n", 0, index) + 1,
            "context": content[context_start:context_end],
        })
        start_at = index + max(1, len(needle))
    return {"success": True, "match_count": len(matches), "matches": matches,
            "metadata": result["metadata"]}
